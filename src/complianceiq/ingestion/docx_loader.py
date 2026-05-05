"""DOCX loader using python-docx.

Walks the document body in order, capturing paragraphs and tables, and
maintains a heading hierarchy (Heading 1/2/3 → 'A > B > C') as section context.
"""

from __future__ import annotations

import logging
import os
from pathlib import Path
from typing import Callable

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentT
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from complianceiq.models import Document, DocumentMetadata, Domain

logger = logging.getLogger(__name__)

DomainInferer = Callable[[str], Domain]


class DOCXLoader:
    """Loads policy DOCX files preserving paragraph order, headings, and tables."""

    def __init__(self, domain_inferer: DomainInferer) -> None:
        self.infer_domain = domain_inferer
        self.failed_files: list[tuple[str, str]] = []

    # ── helpers ───────────────────────────────────────────────
    @staticmethod
    def _is_docx(name: str) -> bool:
        if name.startswith("~$") or name.startswith("."):
            return False
        return name.lower().endswith(".docx")

    @staticmethod
    def _heading_level(style_name: str) -> int | None:
        """Return 1..9 for Heading N / Title; None otherwise."""
        if not style_name:
            return None
        s = style_name.strip().lower()
        if s in ("title", "subtitle"):
            return 1
        if s.startswith("heading"):
            tail = s.replace("heading", "").strip()
            if tail.isdigit():
                return max(1, min(int(tail), 9))
            return 1
        return None

    @staticmethod
    def _iter_block_items(parent: DocxDocumentT):
        """Yield Paragraphs and Tables in document order."""
        body = parent.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, parent)
            elif child.tag == qn("w:tbl"):
                yield Table(child, parent)

    @staticmethod
    def _table_to_text(table: Table) -> str:
        rows: list[str] = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        return "\n".join(rows)

    # ── main API ──────────────────────────────────────────────
    def load_all_docx(self, directory: str | os.PathLike) -> list[Document]:
        directory = Path(directory)
        if not directory.exists():
            logger.warning("DOCX directory does not exist: %s", directory)
            return []

        docx_files = sorted(
            p for p in directory.rglob("*") if p.is_file() and self._is_docx(p.name)
        )
        logger.info("Found %d DOCX files in %s", len(docx_files), directory)

        documents: list[Document] = []
        for path in docx_files:
            documents.extend(self._load_one(path))

        logger.info("DOCX loading complete: %d documents | %d failed",
                    len(documents), len(self.failed_files))
        return documents

    def _load_one(self, path: Path) -> list[Document]:
        documents: list[Document] = []
        domain = self.infer_domain(str(path))

        try:
            doc = DocxDocument(str(path))
        except Exception as e:
            logger.exception("Failed to open DOCX %s", path.name)
            self.failed_files.append((str(path), str(e)))
            return documents

        # Maintain heading stack: heading_stack[i-1] = current heading at level i.
        heading_stack: list[str] = []

        def current_section() -> str:
            return " > ".join(h for h in heading_stack if h)

        try:
            for block in self._iter_block_items(doc):
                if isinstance(block, Paragraph):
                    text = block.text.strip()
                    if not text:
                        continue

                    style_name = block.style.name if block.style else ""
                    level = self._heading_level(style_name)

                    if level is not None:
                        # Update heading stack at the given level, truncate deeper levels.
                        while len(heading_stack) < level:
                            heading_stack.append("")
                        heading_stack[level - 1] = text
                        del heading_stack[level:]
                        # Also persist the heading text itself as a Document
                        # so it can be embedded/searched.
                        documents.append(Document(
                            text=text,
                            metadata=DocumentMetadata(
                                source=str(path),
                                file_name=path.name,
                                doc_type="policy",
                                domain=domain,
                                section=current_section(),
                                block_type="heading",
                            ),
                        ))
                        continue

                    documents.append(Document(
                        text=text,
                        metadata=DocumentMetadata(
                            source=str(path),
                            file_name=path.name,
                            doc_type="policy",
                            domain=domain,
                            section=current_section(),
                            block_type="paragraph",
                        ),
                    ))

                elif isinstance(block, Table):
                    table_text = self._table_to_text(block)
                    if not table_text:
                        continue
                    documents.append(Document(
                        text=table_text,
                        metadata=DocumentMetadata(
                            source=str(path),
                            file_name=path.name,
                            doc_type="policy",
                            domain=domain,
                            section=current_section(),
                            block_type="table",
                        ),
                    ))

        except Exception as e:
            logger.exception("Error while parsing DOCX %s", path.name)
            self.failed_files.append((str(path), str(e)))

        return documents
