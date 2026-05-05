"""Audit-ready DOCX assessment report.

Reads all Week 5 artifacts and produces a single Word document with:
  - Cover + executive summary
  - Per-domain coverage table
  - Top remediation items
  - Top high-severity gaps with evidence
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_ALIGN_VERTICAL

from complianceiq.config import REPORT_DOCX_FILE, ensure_reports_dir
from complianceiq.models import CoverageQuality, Severity
from complianceiq.reporting.loader import (
    load_coverage,
    load_evidence,
    load_gaps,
    load_remediation,
    load_score,
)

logger = logging.getLogger(__name__)

_SEVERITY_COLOR = {
    "high":   RGBColor(0xC0, 0x39, 0x2B),
    "medium": RGBColor(0xE6, 0x7E, 0x22),
    "low":    RGBColor(0x7F, 0x8C, 0x8D),
}


# ── helpers ───────────────────────────────────────────────────
def _add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def _add_kv_row(table, label, value, bold_value=False):
    row = table.add_row().cells
    row[0].text = label
    row[1].text = str(value)
    if bold_value:
        for r in row[1].paragraphs[0].runs:
            r.bold = True


def _styled_cell(cell, text, color: RGBColor | None = None, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    if color is not None:
        run.font.color.rgb = color
    run.bold = bold


# ── sections ──────────────────────────────────────────────────
def _section_cover(doc, score) -> None:
    title = doc.add_heading("ComplianceIQ", level=0)
    sub = doc.add_paragraph("AI-Driven Multi-Agent Regulatory Compliance Assessment")
    sub.runs[0].italic = True
    doc.add_paragraph(f"Generated: {datetime.utcnow().isoformat(timespec='seconds')} UTC")

    if score is None:
        doc.add_paragraph(
            "No compliance_score.json found. Run `python main.py orchestrate` first."
        )
        return

    doc.add_paragraph()
    headline = doc.add_paragraph()
    run = headline.add_run(
        f"Overall weighted coverage: {score.weighted_overall_score:.2f}  "
        f"(unweighted: {score.overall_score:.2f})"
    )
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph(score.risk_summary)


def _section_executive_summary(doc, score) -> None:
    _add_heading(doc, "1. Executive Summary", level=1)
    if score is None:
        doc.add_paragraph("No score data available.")
        return

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"
    _add_kv_row(table, "Total regulatory requirements analyzed", score.total_requirements, bold_value=True)
    _add_kv_row(table, "Fully covered", score.fully_covered)
    _add_kv_row(table, "Partially covered", score.partially_covered)
    _add_kv_row(table, "Not covered", score.not_covered)
    _add_kv_row(table, "High-priority gaps", score.high_priority_gaps, bold_value=True)
    _add_kv_row(table, "Overall coverage (unweighted)", f"{score.overall_score:.2f}")
    _add_kv_row(table, "Overall coverage (severity-weighted)", f"{score.weighted_overall_score:.2f}",
                bold_value=True)


def _section_domain_coverage(doc, score) -> None:
    _add_heading(doc, "2. Per-Domain Coverage", level=1)
    if score is None or not score.by_domain:
        doc.add_paragraph("No per-domain data available.")
        return

    coverage_assessments = {a.domain.value: a for a in load_coverage()}

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Domain", "Weighted", "Total", "Open (H/M/L)", "Weakest areas"]):
        hdr[i].text = h
        for r in hdr[i].paragraphs[0].runs:
            r.bold = True

    for d in score.by_domain:
        row = table.add_row().cells
        row[0].text = d.domain.value
        row[1].text = f"{d.weighted_score:.2f}"
        row[2].text = str(d.total_requirements)
        row[3].text = f"{d.high_severity_open}/{d.medium_severity_open}/{d.low_severity_open}"
        weakest = coverage_assessments.get(d.domain.value)
        row[4].text = ", ".join(weakest.weakest_areas[:3]) if weakest else ""


def _section_remediation(doc, top_n: int = 20) -> None:
    _add_heading(doc, "3. Top Remediation Items", level=1)
    items = load_remediation()
    if not items:
        doc.add_paragraph("No remediation_plan.jsonl found. Run `python main.py orchestrate` "
                          "or `python main.py remediate` first.")
        return

    items = sorted(items, key=lambda i: i.priority_score, reverse=True)[:top_n]

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["#", "Domain", "Severity", "Coverage", "Mandatory action"]):
        hdr[i].text = h
        for r in hdr[i].paragraphs[0].runs:
            r.bold = True

    for idx, item in enumerate(items, start=1):
        row = table.add_row().cells
        row[0].text = str(idx)
        row[1].text = item.domain.value
        sev_color = _SEVERITY_COLOR.get(item.severity.value)
        _styled_cell(row[2], item.severity.value, color=sev_color, bold=True)
        row[3].text = item.coverage.value
        row[4].text = item.mandatory_action[:200]


def _section_high_severity_gaps(doc, max_gaps: int = 10) -> None:
    _add_heading(doc, "4. High-Severity Gaps with Evidence", level=1)
    gaps = load_gaps()
    if not gaps:
        doc.add_paragraph("No gaps.jsonl found.")
        return

    high_open = [
        g for g in gaps
        if g.severity == Severity.HIGH and g.coverage != CoverageQuality.FULLY_COVERED
    ][:max_gaps]
    if not high_open:
        doc.add_paragraph("No high-severity open gaps. Excellent.")
        return

    evidence_by_id = {p.requirement_id: p for p in load_evidence()}

    for i, g in enumerate(high_open, start=1):
        _add_heading(doc, f"4.{i}  {g.regulation_source}", level=2)

        p = doc.add_paragraph()
        p.add_run("Regulatory clause: ").bold = True
        p.add_run(g.regulation_text[:600])

        p = doc.add_paragraph()
        p.add_run("Mandatory action: ").bold = True
        p.add_run(g.mandatory_action)

        p = doc.add_paragraph()
        p.add_run("Coverage: ").bold = True
        p.add_run(g.coverage.value)
        p.add_run("    Severity: ").bold = True
        p.add_run(g.severity.value)
        p.add_run("    Priority: ").bold = True
        p.add_run(g.remediation_priority.value)

        if g.gap_description:
            p = doc.add_paragraph()
            p.add_run("Gap: ").bold = True
            p.add_run(g.gap_description)

        if g.remediation_steps:
            doc.add_paragraph("Remediation steps:", style="List Bullet")
            for s in g.remediation_steps[:5]:
                doc.add_paragraph(s, style="List Bullet 2")

        packet = evidence_by_id.get(g.requirement_id)
        if packet and packet.policy_evidence:
            p = doc.add_paragraph()
            p.add_run("Policy evidence considered:").bold = True
            for ev in packet.policy_evidence[:3]:
                doc.add_paragraph(
                    f"{ev.source_file} ({ev.location}) — {ev.text[:200]}",
                    style="List Bullet",
                )


# ── public API ────────────────────────────────────────────────
def generate_docx_report(out_path: Path = REPORT_DOCX_FILE) -> Path:
    ensure_reports_dir()
    score = load_score()

    doc = Document()
    _section_cover(doc, score)
    doc.add_page_break()
    _section_executive_summary(doc, score)
    _section_domain_coverage(doc, score)
    _section_remediation(doc, top_n=20)
    _section_high_severity_gaps(doc, max_gaps=10)

    doc.save(out_path)
    logger.info("Wrote DOCX report to %s", out_path)
    print(f"Report saved to: {out_path}")
    return out_path
